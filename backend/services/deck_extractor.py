from pathlib import Path
import fitz
from pptx import Presentation

def extract_deck(path: str) -> dict:
    p=Path(path); suffix=p.suffix.lower()
    pages=[]
    if suffix==".pdf":
        doc=fitz.open(path)
        for i,page in enumerate(doc):
            pages.append({"page":i+1,"text":page.get_text("text").strip()})
        doc.close()
    elif suffix==".pptx":
        prs=Presentation(path)
        for i,slide in enumerate(prs.slides):
            texts=[]
            for shape in slide.shapes:
                if hasattr(shape,"text") and shape.text:
                    texts.append(shape.text.strip())
            pages.append({"page":i+1,"text":"\n".join(texts)})
    else:
        raise ValueError("Unsupported deck format.")
    text="\n".join(x["text"] for x in pages)
    return {"file_name":p.name,"file_type":suffix[1:].upper(),"page_count":len(pages),"pages":pages,"text":text}
